import pptx

TOOL_SCHEMA = {
    "type": "function",
    "function": {
        "name": "read_pptx",
        "description": "Reads a PowerPoint file and returns the text.",
        "parameters": {
            "type": "object",
            "properties": {
                "path": {"type": "string", "description": "Path to the PowerPoint file"}
            },
            "required": ["path"]
        }
    }
}
def execute(args):
    prs = pptx.Presentation(args['path'])
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text